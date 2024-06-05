import os
import sys
import openpyxl
import pyodbc
import shutil
from datetime import datetime
import gc
import logging

# Set up logging
# logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def get_database_connection():
    dsn = input("Enter DSN: ")
    user = input("Enter username: ")
    password = input("Enter password: ")
    
    connection_string = f"DSN={dsn};UID={user};PWD={password}"
    # logging.info(f"Connecting to database with DSN: {dsn}")
    conn = pyodbc.connect(connection_string)
    return conn

def create_table(cursor, table_name, columns):
    columns_definition = ", ".join([f"{col} NVARCHAR(MAX)" for col in columns])
    create_table_sql = f"CREATE TABLE {table_name} ({columns_definition});"
    # logging.info(f"Creating table {table_name} with columns: {columns_definition}")
    cursor.execute(create_table_sql)

def insert_data(cursor, table_name, columns, data, batch_size=1000):
    placeholders = ", ".join(["?" for _ in columns])
    columns_joined = ", ".join(columns)
    insert_sql = f"INSERT INTO {table_name} ({columns_joined}) VALUES ({placeholders});"

    # logging.info(f"Inserting data into table {table_name} in batches of {batch_size}")
    for i in range(0, len(data), batch_size):
        batch = data[i:i+batch_size]
        cursor.executemany(insert_sql, batch)
        cursor.connection.commit()  # Commit after each batch

def process_workbook(cursor, workbook_path):
    # logging.info(f"Processing workbook: {workbook_path}")
    workbook = openpyxl.load_workbook(workbook_path)
    try:
        for sheet in workbook.worksheets:
            table_name = f"{sheet.title}_{datetime.now().strftime('%Y%m%d')}"
            columns = [cell.value for cell in sheet[1]]
            data = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]

            create_table(cursor, table_name, columns)
            insert_data(cursor, table_name, columns, data)
    finally:
        workbook.close()  # Ensure the workbook is closed after processing
        del workbook
        gc.collect()

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    input_dir = os.path.join(base_dir, 'Input')
    history_dir = os.path.join(base_dir, 'History')

    conn = get_database_connection()
    cursor = conn.cursor()

    try:
        for file_name in os.listdir(input_dir):
            if file_name.endswith(".xlsx"):
                file_path = os.path.join(input_dir, file_name)
                process_workbook(cursor, file_path)
                shutil.move(file_path, os.path.join(history_dir, file_name))
    finally:
        cursor.close()
        conn.close()

if __name__ == "__main__":
    main()
